"""
Gerador do comunicado de decisão no formato Pirâmide de Minto (.docx).

Peça de decisão, par do relatório IMRaD: resposta primeiro, razões depois,
pedido de ação no fim. Ver docs/guia-metodologico.md (raiz do repo), seção 3.
Deriva do relatório por poda, nunca por resumo: não pode conter achado que não
esteja no IMRaD, e cada porquê aponta para uma seção específica dele.

Cada seção vive em secoes/<pasta>/ com um markdown (nome definido em config.yml,
chave "arquivo") e uma subpasta figuras/. O cabeçalho vem de config.yml, chave
documento.cabecalho (Para, De, Data, Ref.), diferente do cabeçalho de relatório
do gerador IMRaD. O estilo é herdado de modelo_de_escrita.docx.

A seção marcada com "contar_porques: true" tem a quantidade de razões validada
contra a faixa de config.yml (chave "porques"): fora dela, o gerador aborta.

Sintaxe de markdown suportada:

    ## Título              -> subtítulo
    ### Título             -> subtítulo menor (é assim que cada porquê é marcado)
    - item / * item        -> lista com marcadores
    1. item                -> lista numerada
    **negrito**, *itálico*, `monoespaçado`
    ![legenda](figuras/x.png)
    | col | col |          -> tabela (com a linha |---|---| abaixo do cabeçalho)
    <!-- orientação -->    -> comentário; nunca aparece no .docx gerado
"""

from pathlib import Path
import re
import sys

import yaml
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

try:
    from PIL import Image
    TEM_PIL = True
except ImportError:
    TEM_PIL = False

# --- Constantes de estilo (extraídas do modelo_de_escrita.docx) ---------------
RAIZ = Path(__file__).resolve().parent
AZUL = RGBColor(0x00, 0x20, 0x60)
CINZA_LEGENDA = RGBColor(0x59, 0x59, 0x59)
VERMELHO_ERRO = RGBColor(0xC0, 0x00, 0x00)
COR_TEXTO_TABELA = RGBColor(0x1A, 0x1A, 0x2E)
FILL_LABEL = "E6ECF7"
FONTE_TITULOS = "Calibri"
TAM_TITULO = 16
TAM_SUBTITULO = 11
TAM_SECAO = 15
TAM_TABELA = 9.5
LARGURA_CONTEUDO_CM = 16.0  # largura útil aproximada da página (A4 com margens)

# Estilo do texto normal (parágrafos e listas) — sobrescrito por config.yml (chave "estilo").
FONTE_CORPO = "Calibri"
FONTE_MONO = "Consolas"  # usada em trechos entre crases (`nome_da_tabela`)
TAM_CORPO = 11
COR_CORPO = RGBColor(0x00, 0x20, 0x60)

PADRAO_INLINE = re.compile(r"(`[^`]+`|\*\*.+?\*\*|\*.+?\*)")
PADRAO_IMAGEM = re.compile(r"^!\[(.*?)\]\((.*?)\)\s*$")
PADRAO_HEADING = re.compile(r"^(#{2,3})\s+(.*)$")
PADRAO_BULLET = re.compile(r"^[-*]\s+(.*)$")
PADRAO_NUMERO = re.compile(r"^\d+\.\s+(.*)$")
PADRAO_SEPARADOR_TABELA = re.compile(r"^[\s\-:|]+$")
PADRAO_COMENTARIO = re.compile(r"<!--.*?-->", re.DOTALL)


def sem_comentarios(texto):
    """Remove blocos <!-- --> para checar se sobra conteúdo real (usado na seção opcional)."""
    return PADRAO_COMENTARIO.sub("", texto).strip()


# --- Utilidades de baixo nível ------------------------------------------------
def carregar_config():
    with open(RAIZ / "config.yml", encoding="utf-8") as f:
        return yaml.safe_load(f)


def aplicar_estilo_config(cfg):
    """Sobrescreve fonte/tamanho/cor do texto normal a partir de config.yml (chave 'estilo')."""
    global FONTE_CORPO, TAM_CORPO, COR_CORPO
    estilo = cfg.get("estilo", {})
    FONTE_CORPO = estilo.get("fonte", FONTE_CORPO)
    TAM_CORPO = estilo.get("tamanho_corpo", TAM_CORPO)
    cor_hex = estilo.get("cor_texto")
    if cor_hex:
        COR_CORPO = RGBColor.from_string(cor_hex)


def limpar_corpo(doc):
    """Remove todo o conteúdo do corpo, preservando o sectPr (config de página)."""
    body = doc.element.body
    for filho in list(body):
        if filho.tag == qn("w:sectPr"):
            continue
        body.remove(filho)


def mover_sectpr_para_fim(doc):
    """Garante que o sectPr seja o último elemento do corpo (docx válido)."""
    body = doc.element.body
    for sect in body.findall(qn("w:sectPr")):
        body.remove(sect)
        body.append(sect)


def sombrear_celula(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def bordas_tabela(tab, cor="BFBFBF"):
    """Aplica bordas em toda a tabela via XML, sem depender de estilo do template."""
    tblPr = tab._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{lado}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), cor)
        borders.append(e)
    tblPr.append(borders)


def estilo_run(run, tam=None, cor=None, negrito=None, italico=None, fonte=FONTE_TITULOS):
    if fonte:
        run.font.name = fonte
    if tam is not None:
        run.font.size = Pt(tam)
    if cor is not None:
        run.font.color.rgb = cor
    if negrito is not None:
        run.bold = negrito
    if italico is not None:
        run.italic = italico


def adicionar_inline(par, texto, tam=None, cor=None, fonte=FONTE_TITULOS):
    """Escreve texto em um parágrafo interpretando **negrito**, *itálico* e `monoespaçado`."""
    for token in PADRAO_INLINE.split(texto):
        if not token:
            continue
        fonte_token = fonte
        if token.startswith("`") and token.endswith("`"):
            run = par.add_run(token[1:-1])
            fonte_token = FONTE_MONO
        elif token.startswith("**") and token.endswith("**"):
            run = par.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("*") and token.endswith("*"):
            run = par.add_run(token[1:-1])
            run.italic = True
        else:
            run = par.add_run(token)
        estilo_run(run, tam=tam, cor=cor, fonte=fonte_token)


# --- Blocos do documento ------------------------------------------------------
def add_titulo(doc, titulo, subtitulo):
    p = doc.add_paragraph()
    estilo_run(p.add_run(titulo), tam=TAM_TITULO, cor=AZUL, negrito=True)
    if subtitulo:
        ps = doc.add_paragraph()
        estilo_run(ps.add_run(subtitulo), tam=TAM_SUBTITULO, cor=AZUL)


def add_tabela_meta(doc, meta):
    """Monta a tabela de cabeçalho.

    Numa peça Minto o cabeçalho é o de um comunicado (Para, De, Data, Ref.),
    não o de um relatório. Por isso ele vem de config.yml, chave
    documento.cabecalho, uma lista de pares rotulo/valor. Sem essa chave, cai
    no cabeçalho de relatório do gerador IMRaD.
    """
    cabecalho = meta.get("cabecalho")
    if cabecalho:
        linhas = [(item.get("rotulo", ""), item.get("valor", "")) for item in cabecalho]
    else:
        linhas = [
            ("Data da elaboração", meta.get("data", "")),
            ("Janela temporal", meta.get("janela_temporal", "")),
            ("Base de dados", meta.get("base_de_dados", "")),
            ("Contexto", meta.get("contexto", "")),
        ]
    tab = doc.add_table(rows=0, cols=2)
    tab.autofit = False
    bordas_tabela(tab)
    for rotulo, valor in linhas:
        celulas = tab.add_row().cells
        celulas[0].text = ""
        sombrear_celula(celulas[0], FILL_LABEL)
        estilo_run(celulas[0].paragraphs[0].add_run(rotulo), tam=TAM_TABELA, cor=AZUL, negrito=True)
        celulas[1].text = ""
        valores = valor if isinstance(valor, list) else [valor]
        for i, v in enumerate(valores):
            par = celulas[1].paragraphs[0] if i == 0 else celulas[1].add_paragraph()
            adicionar_inline(par, str(v), tam=TAM_TABELA, cor=COR_TEXTO_TABELA)
    for linha in tab.rows:
        linha.cells[0].width = Cm(3.8)
        linha.cells[1].width = Cm(12.2)
    doc.add_paragraph()


def eh_linha_separadora_tabela(linha):
    """Detecta a linha `|---|---|` que separa cabeçalho e dados numa tabela."""
    return bool(linha) and "-" in linha and "|" in linha and bool(PADRAO_SEPARADOR_TABELA.match(linha))


def celulas_tabela(linha):
    """Quebra uma linha `| a | b |` em ["a", "b"]."""
    linha = linha.strip()
    if linha.startswith("|"):
        linha = linha[1:]
    if linha.endswith("|"):
        linha = linha[:-1]
    return [c.strip() for c in linha.split("|")]


def add_tabela_markdown(doc, cabecalho, dados):
    """Renderiza uma tabela markdown (cabeçalho + linhas) como tabela do Word."""
    n_cols = len(cabecalho)
    tab = doc.add_table(rows=0, cols=n_cols)
    tab.autofit = False
    bordas_tabela(tab)

    celulas_cab = tab.add_row().cells
    for i, texto in enumerate(cabecalho):
        sombrear_celula(celulas_cab[i], FILL_LABEL)
        par = celulas_cab[i].paragraphs[0]
        adicionar_inline(par, texto, tam=TAM_TABELA, cor=AZUL)
        for run in par.runs:
            run.font.bold = True

    for linha in dados:
        celulas = tab.add_row().cells
        for i in range(n_cols):
            texto = linha[i] if i < len(linha) else ""
            adicionar_inline(celulas[i].paragraphs[0], texto, tam=TAM_TABELA, cor=COR_TEXTO_TABELA)

    largura_col = Cm(LARGURA_CONTEUDO_CM / n_cols)
    for linha in tab.rows:
        for cell in linha.cells:
            cell.width = largura_col

    doc.add_paragraph()


def add_heading_secao(doc, texto):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    estilo_run(p.add_run(texto), tam=TAM_SECAO, cor=AZUL, negrito=True)


def largura_imagem(caminho):
    """Largura da imagem sem estourar a página; encolhe só se for larga demais."""
    if TEM_PIL:
        try:
            with Image.open(caminho) as im:
                dpi = im.info.get("dpi", (96, 96))[0] or 96
                larg_cm = im.width / dpi * 2.54
                if larg_cm < LARGURA_CONTEUDO_CM:
                    return Cm(larg_cm)
        except Exception:
            pass
    return Cm(LARGURA_CONTEUDO_CM)


def add_item_lista(doc, texto, marcador):
    """Item de lista com recuo pendente (não depende de estilos do template)."""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Pt(18)
    pf.first_line_indent = Pt(-18)
    pf.space_after = Pt(2)
    pf.tab_stops.add_tab_stop(Pt(18))
    estilo_run(p.add_run(marcador + "\t"), tam=TAM_CORPO, cor=COR_CORPO, fonte=FONTE_CORPO)
    adicionar_inline(p, texto, tam=TAM_CORPO, cor=COR_CORPO, fonte=FONTE_CORPO)


def add_imagem(doc, caminho, legenda):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(caminho), width=largura_imagem(caminho))
    if legenda:
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        estilo_run(cap.add_run(legenda), tam=9, cor=CINZA_LEGENDA, italico=True)


# --- Renderização do markdown -------------------------------------------------
def renderizar_markdown(doc, texto_md, pasta_secao):
    linhas = texto_md.splitlines()
    buffer = []
    contador_num = 0
    i = 0
    n = len(linhas)

    def descarregar():
        if buffer:
            adicionar_inline(doc.add_paragraph(), " ".join(buffer).strip(), tam=TAM_CORPO, cor=COR_CORPO, fonte=FONTE_CORPO)
            buffer.clear()

    while i < n:
        linha = linhas[i]
        crua = linha.strip()

        # Linha em branco e comentário não zeram a numeração: itens de lista
        # separados por linha vazia (ou por uma orientação <!-- -->) continuam
        # a mesma lista, como em markdown. Quem zera é conteúdo de outro tipo
        # (parágrafo, título, tabela, figura, bullet), tratado abaixo.
        if crua == "":
            descarregar()
            i += 1
            continue

        if crua.startswith("<!--"):
            descarregar()
            if "-->" not in crua:
                i += 1
                while i < n and "-->" not in linhas[i]:
                    i += 1
            i += 1
            continue

        if "|" in crua and i + 1 < n and eh_linha_separadora_tabela(linhas[i + 1].strip()):
            descarregar()
            contador_num = 0
            cabecalho = celulas_tabela(crua)
            i += 2
            dados = []
            while i < n and linhas[i].strip() and "|" in linhas[i]:
                dados.append(celulas_tabela(linhas[i]))
                i += 1
            add_tabela_markdown(doc, cabecalho, dados)
            continue

        m = PADRAO_IMAGEM.match(crua)
        if m:
            descarregar()
            contador_num = 0
            legenda, rel = m.group(1), m.group(2)
            caminho = (pasta_secao / rel).resolve()
            if caminho.exists():
                add_imagem(doc, caminho, legenda)
            else:
                p = doc.add_paragraph()
                estilo_run(p.add_run(f"[FIGURA NÃO ENCONTRADA: {rel}]"), cor=VERMELHO_ERRO)
            i += 1
            continue

        m = PADRAO_HEADING.match(linha)
        if m:
            descarregar()
            contador_num = 0
            tam = 13 if len(m.group(1)) == 2 else 12
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            estilo_run(p.add_run(m.group(2).strip()), tam=tam, cor=AZUL, negrito=True)
            i += 1
            continue

        m = PADRAO_BULLET.match(crua)
        if m:
            descarregar()
            contador_num = 0
            add_item_lista(doc, m.group(1), "•")
            i += 1
            continue

        m = PADRAO_NUMERO.match(crua)
        if m:
            descarregar()
            contador_num += 1
            add_item_lista(doc, m.group(1), f"{contador_num}.")
            i += 1
            continue

        contador_num = 0
        buffer.append(crua)
        i += 1

    descarregar()


# --- Orquestração -------------------------------------------------------------
def validar_porques(sec, texto, regra):
    """Confere que a seção dos porquês tem uma quantidade razoável de razões.

    A pirâmide de Minto pede razões organizadas de forma MECE (mutuamente
    exclusivas, coletivamente exaustivas). Menos de 3 raramente sustenta a
    recomendação; mais de 10 vira lista, e o leitor perde a hierarquia. Cada
    porquê é um '### Título' no markdown. Fora da faixa, aborta: é erro de
    conteúdo, não de formatação, e passar batido produz uma peça fraca.
    """
    minimo = regra.get("minimo", 3)
    maximo = regra.get("maximo", 10)
    achados = [
        linha for linha in sem_comentarios(texto).splitlines()
        if re.match(r"^###\s+\S", linha)
    ]
    n = len(achados)
    if not minimo <= n <= maximo:
        print(
            f"[erro] seção '{sec['pasta']}' tem {n} porquê(s); o config.yml exige "
            f"entre {minimo} e {maximo}. Cada porquê é um '### Título' no markdown.",
            file=sys.stderr,
        )
        sys.exit(1)
    print(f"[ok] {n} porquês (faixa permitida: {minimo} a {maximo})")


def main():
    cfg = carregar_config()
    aplicar_estilo_config(cfg)

    base = RAIZ / cfg.get("modelo_base", "modelo_de_escrita.docx")
    if base.exists():
        doc = Document(str(base))
        limpar_corpo(doc)
    else:
        print(f"[aviso] modelo base não encontrado ({base}); usando documento em branco.")
        doc = Document()

    meta = cfg.get("documento", {})
    add_titulo(doc, meta.get("titulo", "[Título]"), meta.get("subtitulo", ""))
    add_tabela_meta(doc, meta)

    numero = 0
    for sec in cfg.get("secoes", []):
        pasta = RAIZ / "secoes" / sec["pasta"]
        nome_arquivo = sec.get("arquivo")
        if not nome_arquivo:
            print(f"[erro] seção '{sec['pasta']}' sem chave 'arquivo' no config.yml", file=sys.stderr)
            sys.exit(1)
        md_path = pasta / nome_arquivo
        texto = md_path.read_text(encoding="utf-8").strip() if md_path.exists() else ""
        if sec.get("opcional") and not sem_comentarios(texto):
            print(f"[pulado] {sec['pasta']} (opcional e vazio)")
            continue
        if sec.get("numerar", True):
            numero += 1
            titulo_secao = f"{numero}. {sec['titulo']}"
        else:
            titulo_secao = sec["titulo"]
        if sec.get("contar_porques"):
            validar_porques(sec, texto, cfg.get("porques", {}))
        add_heading_secao(doc, titulo_secao)
        renderizar_markdown(doc, texto, pasta)

    mover_sectpr_para_fim(doc)
    saida = RAIZ / "saida"
    saida.mkdir(exist_ok=True)
    destino = saida / cfg.get("arquivo_saida", "analise_gerada.docx")
    doc.save(str(destino))
    print(f"OK -> {destino}")


if __name__ == "__main__":
    try:
        main()
    except FileNotFoundError as e:
        print(f"[erro] arquivo não encontrado: {e}", file=sys.stderr)
        sys.exit(1)
